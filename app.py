# app.py
# Author: MMussie + M365 Copilot
# Version: 2025.11.19-p1
# Two-stage sampling app with PPS village selection (balanced MOS),
# systematic WOR within cells, configurable replacement rate, printing-ready Excel export,
# and robust sheet detection for 'Frame_Villages' variants (e.g., 'Frame Villages').

import io
import math
import hashlib
from datetime import datetime
import numpy as np
import pandas as pd
import streamlit as st

st.set_page_config(page_title="Two-Stage PPS Sampling (WFP Somali Region)", layout="wide")

# ---------------- HELPERS ----------------
def stable_data_seed(df_villages: pd.DataFrame, df_hhs: pd.DataFrame) -> int:
    parts = []
    if df_villages is not None and not df_villages.empty:
        cols_v = ["Kebele", "Village", "E_count", "I_count"]
        cols_v = [c for c in cols_v if c in df_villages.columns]
        parts.extend(["|".join(map(str,row)) for row in df_villages[cols_v].values])
    if df_hhs is not None and not df_hhs.empty:
        cols_h = ["Kebele", "Village", "Group", "HH_ID"]
        cols_h = [c for c in cols_h if c in df_hhs.columns]
        parts.extend(["|".join(map(str,row)) for row in df_hhs[cols_h].values])
    h = hashlib.sha256("\n".join(parts).encode()).hexdigest()[:8]
    return int(h,16)

def normalize_groups(df):
    gmap = {"e":"E","eligible":"E","1":"E","i":"I","ineligible":"I","0":"I"}
    return df["Group"].astype(str).str.strip().str.lower().map(gmap).fillna(df["Group"])

def pick_rng(seed: int, salt: str) -> np.random.RandomState:
    h = hashlib.sha256(f"{seed}-{salt}".encode()).hexdigest()[:8]
    return np.random.RandomState(int(h,16))

def ensure_int(x, default=0):
    try: return int(x)
    except: return default

# ---------------- PPS STAGE ----------------
def compute_mos_cols(df_villages: pd.DataFrame, mos_method: str) -> pd.DataFrame:
    df = df_villages.copy()
    rename_map = {}
    if "Eligible" in df.columns and "E_count" not in df.columns: rename_map["Eligible"] = "E_count"
    if "Ineligible" in df.columns and "I_count" not in df.columns: rename_map["Ineligible"] = "I_count"
    if "E" in df.columns and "E_count" not in df.columns: rename_map["E"] = "E_count"
    if "I" in df.columns and "I_count" not in df.columns: rename_map["I"] = "I_count"
    df = df.rename(columns=rename_map)

    df["E_count"] = df["E_count"].apply(ensure_int)
    df["I_count"] = df["I_count"].apply(ensure_int)
    df["minEI"] = df[["E_count","I_count"]].min(axis=1)
    df["sqrtEI"] = np.sqrt(df["E_count"].clip(lower=0) * df["I_count"].clip(lower=0))
    df["totalEI"] = df["E_count"] + df["I_count"]

    df["MOS_raw"] = df["minEI"] if mos_method == "Balanced (min(E,I))" else df["sqrtEI"]
    df["MOS"] = df["MOS_raw"]
    for keb, grp in df.groupby("Kebele"):
        if (grp["minEI"] == 0).all():
            df.loc[df["Kebele"]==keb, "MOS"] = df.loc[df["Kebele"]==keb, "totalEI"]
    df["MOS"] = df["MOS"].fillna(0).clip(lower=0)
    return df

def pps_without_replacement(df_k: pd.DataFrame, max_villages: int, rng: np.random.RandomState):
    candidates = df_k[df_k["MOS"]>0].copy()
    if candidates.empty:
        candidates = df_k.copy()
        if (candidates["E_count"] + candidates["I_count"] <= 0).all():
            return df_k.iloc[0:0]
    max_pick = min(max_villages, len(candidates))
    selected_idxs = []
    remaining = candidates.index.tolist()
    for order in range(1, max_pick+1):
        sub = df_k.loc[remaining]
        weights = sub["MOS"].to_numpy(dtype=float)
        probs = weights/weights.sum() if weights.sum()>0 else np.ones_like(weights)/len(weights)
        choice = rng.choice(sub.index.to_numpy(), p=probs)
        selected_idxs.append(choice)
        remaining.remove(choice)
    sel = df_k.loc[selected_idxs].copy()
    sel["Selected_Order"] = range(1,len(sel)+1)
    total_mos = df_k["MOS"].sum()
    sel["Indicative_Prob"] = sel["MOS"]/total_mos if total_mos>0 else 0
    return sel

# ---------------- ALLOCATION ----------------
def allocate_within_kebele(kebele: str, sel_villages: pd.DataFrame, counts_k: pd.DataFrame,
                           group_targets_per_village: dict, skip_ineligible_if_under_30: bool=True):
    cap = (counts_k.pivot_table(index=["Kebele","Village"], columns="Group", values="N", aggfunc="sum", fill_value=0).reset_index())
    for g in ["E","I"]:
        if g not in cap.columns: cap[g] = 0
    sel_names = [v for v in sel_villages["Village"].tolist()]
    cap = cap[cap["Village"].isin(sel_names)].copy()

    total_I = int(cap["I"].sum())
    use_I = not (skip_ineligible_if_under_30 and total_I < 30)

    alloc_rows = []
    for v in sel_names:
        for g in ["E","I"]:
            target = 0 if (g=="I" and not use_I) else group_targets_per_village[g]
            available = int(cap.loc[cap["Village"]==v, g].sum())
            alloc_rows.append([kebele, v, g, target, available])
    alloc = pd.DataFrame(alloc_rows, columns=["Kebele","Village","Group","Target_Primaries","Available"])

    if not use_I:
        capable = cap.loc[cap["E"]>=30, "Village"].tolist()
        preferred = [v for v in sel_villages.sort_values("Selected_Order")["Village"] if v in capable]
        if preferred:
            keep = preferred[0]
            alloc.loc[alloc["Village"]!=keep, "Target_Primaries"] = 0
            alloc.loc[(alloc["Village"]==keep) & (alloc["Group"]=="E"), "Target_Primaries"] = 30

    for g in (["E","I"] if use_I else ["E"]):
        desired = 30
        mask_g = alloc["Group"]==g
        alloc.loc[mask_g, "Target_Primaries"] = np.minimum(alloc.loc[mask_g, "Target_Primaries"], alloc.loc[mask_g, "Available"])
        current = int(alloc.loc[mask_g, "Target_Primaries"].sum())
        if current < desired:
            shortage = desired - current
            for v in sel_villages.sort_values("Selected_Order")["Village"]:
                mask = (alloc["Village"]==v) & mask_g
                spare = int(alloc.loc[mask, "Available"].values[0] - alloc.loc[mask, "Target_Primaries"].values[0])
                if spare<=0: continue
                add = min(spare, shortage)
                alloc.loc[mask, "Target_Primaries"] += add
                shortage -= add
                if shortage<=0: break
        elif current > desired:
            excess = current - desired
            for v in sel_villages.sort_values("Selected_Order", ascending=False)["Village"]:
                mask = (alloc["Village"]==v) & mask_g
                can_cut = int(alloc.loc[mask, "Target_Primaries"].values[0])
                cut = min(can_cut, excess)
                alloc.loc[mask, "Target_Primaries"] -= cut
                excess -= cut
                if excess<=0: break
    alloc["Target_Primaries"] = np.minimum(alloc["Target_Primaries"], alloc["Available"])
    return alloc

# ---------------- SYSTEMATIC SAMPLING & REPLACEMENTS --------------
def systematic_wor(df_cell: pd.DataFrame, n: int, rng: np.random.RandomState, sort_cols=None):
    N=len(df_cell)
    if n<=0 or N<=0: return df_cell.iloc[0:0].copy(), None, 0
    if sort_cols is None or not set(sort_cols).issubset(df_cell.columns):
        df_cell = df_cell.sort_values(by=df_cell.columns.tolist()).reset_index(drop=True)
    else:
        df_cell = df_cell.sort_values(by=sort_cols).reset_index(drop=True)
    k = N/n
    r = rng.uniform(0,k)
    picks = [min(int(r + j*k), N-1) for j in range(n)]
    sel = df_cell.iloc[picks].copy()
    return sel, r, k

def sample_cell(df_hhs, kebele, village, group, n_prim, base_seed, replacement_rate=0.25):
    key = f"{kebele}|{village}|{group}"
    rng = pick_rng(base_seed, f"cell-{key}")
    cell = df_hhs[(df_hhs["Kebele"]==kebele) & (df_hhs["Village"]==village) & (df_hhs["Group"]==group)].copy()
    N = len(cell)
    n_prim = min(int(n_prim), N)
    prim, rstart, step = systematic_wor(cell, n_prim, rng)
    remainder = cell[~cell["HH_ID"].isin(prim["HH_ID"])].copy()
    rate = max(0.0, min(1.0, float(replacement_rate)))
    n_repl = int(math.ceil(rate * n_prim)) if n_prim>0 else 0
    n_repl = min(n_repl, len(remainder))
    repl_rng = pick_rng(base_seed, f"repl-{key}")
    repl, rstart_repl, step_repl = systematic_wor(remainder, n_repl, repl_rng)

    prim["Kebele"] = kebele; prim["Village"] = village; prim["Group"] = group
    prim["Sample_Type"] = "Primary"; prim["RandomStart"] = rstart; prim["Interval"] = step
    repl["Kebele"] = kebele; repl["Village"] = village; repl["Group"] = group
    repl["Sample_Type"] = "Replacement"; repl["RandomStart"] = rstart_repl; repl["Interval"] = step_repl
    repl = repl.reset_index(drop=True); repl["Replacement_Order"] = repl.index + 1
    return prim, repl

# ---------------- EXCEL EXPORT --------------
def autosize_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            v = str(cell.value) if cell.value is not None else ""
            max_len = max(max_len, len(v))
        ws.column_dimensions[col_letter].width = min(max_len+2, 40)

def style_header_row(ws, row_idx: int):
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(border_style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[row_idx]:
        cell.font = bold; cell.fill = fill; cell.alignment = center; cell.border = border

def style_whole_sheet(ws):
    if ws.max_row >= 1: style_header_row(ws, 1)
    autosize_columns(ws)
    from openpyxl.styles import Border, Side
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for c in row: c.border = border

def build_printable_workbook(xl_writer, df_prim, df_repl, kebeles, pass_cols, include_moda, replacement_rate_pct):
    from openpyxl.utils.dataframe import dataframe_to_rows
    book = xl_writer.book
    for keb in kebeles:
        prim_k = df_prim[df_prim["Kebele"]==keb].copy()
        repl_k = df_repl[df_repl["Kebele"]==keb].copy()
        base_cols = ["Kebele","Village","Group","HH_ID"]
        cols = base_cols + [c for c in pass_cols if c not in base_cols]
        if include_moda:
            for c in ["Attempt","Contact_Status","Outcome","Enumerator","Notes"]:
                if c not in cols: cols.append(c)
        prim_cols = [c for c in cols if c in prim_k.columns] + ["Sample_Type"]
        repl_cols = [c for c in cols if c in repl_k.columns] + ["Sample_Type","Replacement_Order"]
        prim_k = prim_k.reindex(columns=prim_cols).sort_values(["Village","Group"])
        repl_k = repl_k.reindex(columns=repl_cols).sort_values(["Village","Group","Replacement_Order"])
        ws = book.create_sheet(title=str(keb)[:31])
        ws.append([f"Kebele: {keb} – Primaries"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(4, len(prim_cols)))
        ws.append([])
        start_prim_header = ws.max_row + 1
        for row in dataframe_to_rows(prim_k, index=False, header=True): ws.append(row)
        style_header_row(ws, start_prim_header); autosize_columns(ws)
        ws.append([])
        ws.append([f"Kebele: {keb} – Replacements (rate: {replacement_rate_pct:.0f}%)"])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=max(4, len(repl_cols)))
        ws.append([])
        if not repl_k.empty:
            start_repl_header = ws.max_row + 1
            for row in dataframe_to_rows(repl_k, index=False, header=True): ws.append(row)
            style_header_row(ws, start_repl_header); autosize_columns(ws)
        else:
            ws.append(["(No replacements)"]); autosize_columns(ws)
        ws.page_setup.orientation = "portrait"
        ws.page_margins.left = 0.3; ws.page_margins.right = 0.3
        ws.page_margins.top = 0.5; ws.page_margins.bottom = 0.5

def export_to_excel(df_prim, df_repl, df_sel_villages, df_alloc, settings, include_moda, printable, replacement_rate_pct):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_prim.to_excel(writer, index=False, sheet_name="Primaries")
        df_repl.to_excel(writer, index=False, sheet_name="Replacements")
        df_sel_villages.to_excel(writer, index=False, sheet_name="Selected_Villages")
        df_alloc.to_excel(writer, index=False, sheet_name="Allocation_Summary")
        pd.DataFrame([settings]).to_excel(writer, index=False, sheet_name="Settings_Log")
        book = writer.book
        for s in ["Primaries","Replacements","Selected_Villages","Allocation_Summary","Settings_Log"]:
            ws = book[s]; style_whole_sheet(ws)
        if printable:
            kebeles = df_alloc["Kebele"].dropna().unique().tolist()
            base_keep = {"Kebele","Village","Group","HH_ID","Sample_Type","RandomStart","Interval","Replacement_Order"}
            pass_cols_prim = [c for c in df_prim.columns if c not in base_keep]
            pass_cols_repl = [c for c in df_repl.columns if c not in base_keep]
            pass_cols = []
            for c in pass_cols_prim + pass_cols_repl:
                if c not in pass_cols and c not in base_keep: pass_cols.append(c)
            build_printable_workbook(writer, df_prim, df_repl, kebeles, pass_cols, include_moda, replacement_rate_pct)
    output.seek(0)
    return output

# ---------------- UI ----------------
st.title("Two-Stage PPS Sampling")
st.caption("Stage 1: PPS village selection per kebele; Stage 2: Systematic WOR within cells; configurable replacements; printing-ready Excel export.")

with st.sidebar:
    st.header("Settings")
    mos_method = st.radio("Balanced MOS method", ["Balanced (min(E,I))", "Balanced (sqrt(E×I))"], index=0)
    max_villages = st.number_input("Max villages per kebele", min_value=1, max_value=10, value=2, step=1)
    st.markdown("**Random Seed (Reproducibility)**")
    seed_mode = st.selectbox("Seed mode", ["Data-derived (best)", "Fixed (20251031)", "Manual"], index=0)
    manual_seed = st.text_input("Manual seed (integer)", value="")
    replacement_rate_pct = st.number_input("Replacement rate (%)", min_value=0.0, max_value=100.0, value=25.0, step=1.0)
    include_moda = st.checkbox("Include MoDa tracking fields", value=True)
    printable = st.checkbox("Create field-team printable workbook (one sheet per kebele)", value=True)
    st.markdown("---")
    pass_cols_user = st.text_input("Pass-through columns (comma-separated)", value="Head_Name,Phone,Address")

# --- Robust sheet name finder ---
def _find_sheet_name(xls, preferred: str, aliases: list):
    def norm(s: str) -> str:
        return s.strip().lower().replace(" ", "").replace("_", "").replace("-", "")
    target_norms = {norm(preferred)} | {norm(a) for a in aliases}
    for sn in xls.sheet_names:
        if norm(sn) in target_norms:
            return sn
    return None

file = st.file_uploader("Upload Excel with sheets 'Frame_Villages' and 'Frame_HHs' (.xlsx)", type=["xlsx"])
run_btn = st.button("▶️ Run Sampling", type="primary", disabled=(file is None))

if run_btn and file is not None:
    try:
        # Villages sheet: canonical then robust fallbacks
        try:
            df_villages = pd.read_excel(file, sheet_name="Frame_Villages", engine="openpyxl")
        except Exception:
            xls = pd.ExcelFile(file, engine="openpyxl")
            preferred = "Frame_Villages"
            aliases = [
                "Frame Villages",   # space
                "Frame-Villages",   # hyphen
                "Villages",         # short
                "Frame_Village",    # singular
                "frame_villages"    # lowercase canonical
            ]
            alt_villages = _find_sheet_name(xls, preferred=preferred, aliases=aliases)
            if alt_villages is None:
                st.error("Unable to find 'Frame_Villages' sheet (accepted variants: 'Frame Villages', 'Frame-Villages', 'Villages', 'Frame_Village').")
                st.stop()
            df_villages = pd.read_excel(file, sheet_name=alt_villages, engine="openpyxl")

        # HHs sheet (existing fallbacks)
        try:
            df_hhs = pd.read_excel(file, sheet_name="Frame_HHs", engine="openpyxl")
        except Exception:
            xls = pd.ExcelFile(file, engine="openpyxl")
            alt_hhs = None
            for sn in xls.sheet_names:
                if sn.strip().lower() in {"frame_hhs", "households", "hh_frame", "frame_households"}:
                    alt_hhs = sn; break
            if alt_hhs is None:
                st.error("Unable to find 'Frame_HHs' sheet."); st.stop()
            df_hhs = pd.read_excel(file, sheet_name=alt_hhs, engine="openpyxl")

        # Validate HHs
        needed_h = {"Kebele","Village","Group","HH_ID"}
        if not needed_h.issubset(df_hhs.columns):
            st.error(f"'Frame_HHs' must include columns: {sorted(needed_h)}"); st.stop()
        df_hhs["Group"] = normalize_groups(df_hhs)

        # Pass-through cols
        pass_cols = [c.strip() for c in pass_cols_user.split(",") if c.strip()]
        keep_cols_hhs = ["Kebele","Village","Group","HH_ID"] + [c for c in pass_cols if c in df_hhs.columns]
        df_hhs = df_hhs[keep_cols_hhs].copy()

        # Seed
        if seed_mode == "Data-derived (best)": seed = stable_data_seed(df_villages, df_hhs)
        elif seed_mode == "Fixed (20251031)": seed = 20251031
        else:
            try: seed = int(manual_seed)
            except: st.error("Manual seed must be an integer."); st.stop()

        # MOS
        df_mos = compute_mos_cols(df_villages, mos_method)
        if not {"Kebele","Village","E_count","I_count"}.issubset(df_mos.columns):
            st.error("'Frame_Villages' must have Kebele, Village, and E_count/I_count (or mappable names like Eligible/Ineligible or E/I)."); st.stop()

        counts = (df_hhs.groupby(["Kebele","Village","Group"], as_index=False).agg(N=("HH_ID","count")))
        sel_records, alloc_records = [], []
        for kebele, grp in df_mos.groupby("Kebele"):
            rng_k = pick_rng(seed, f"kebele-{kebele}")
            selected = pps_without_replacement(grp, max_villages, rng_k)
            if selected.empty: continue
            sel_records.append(selected.assign(Kebele=kebele))
            counts_k = counts[counts["Kebele"]==kebele]
            alloc_k = allocate_within_kebele(kebele, selected, counts_k, {"E":15, "I":15}, skip_ineligible_if_under_30=True)
            alloc_records.append(alloc_k)
        if not sel_records:
            st.warning("No villages selected (check MOS and counts)."); st.stop()

        df_sel = pd.concat(sel_records, ignore_index=True)
        df_alloc = pd.concat(alloc_records, ignore_index=True)

        prim_list, repl_list = [] , []
        for _, row in df_alloc.iterrows():
            kebele = row["Kebele"]; village = row["Village"]; group = row["Group"]
            n_prim = int(row["Target_Primaries"])
            if n_prim <= 0: continue
            prim, repl = sample_cell(df_hhs, kebele, village, group, n_prim, base_seed=seed, replacement_rate=(replacement_rate_pct/100.0))
            prim_list.append(prim)
            if not repl.empty: repl_list.append(repl)
        df_prim = pd.concat(prim_list, ignore_index=True) if prim_list else pd.DataFrame(columns=["Kebele","Village","Group","HH_ID","Sample_Type","RandomStart","Interval"])
        df_repl = pd.concat(repl_list, ignore_index=True) if repl_list else pd.DataFrame(columns=["Kebele","Village","Group","HH_ID","Sample_Type","RandomStart","Interval","Replacement_Order"])

        if include_moda:
            for df_ in [df_prim, df_repl]:
                for c in ["Attempt","Contact_Status","Outcome","Enumerator","Notes"]:
                    if c not in df_.columns: df_[c] = ""

        keep_sv = ["Kebele","Village","E_count","I_count","minEI","sqrtEI","totalEI","MOS","Selected_Order","Indicative_Prob"]
        df_sel_villages = (df_sel[["Kebele","Village"]].merge(df_mos[keep_sv], on=["Kebele","Village"], how="left").drop_duplicates(subset=["Kebele","Village"]))

        settings = {
            "Run_Timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "MOS_Method": mos_method,
            "Max_Villages_per_Kebele": int(max_villages),
            "Seed_Mode": seed_mode,
            "Seed_Value": int(seed),
            "Replacement_Rate": f"{replacement_rate_pct}%",
            "Stage2_Method": "Systematic WOR with random start",
            "Include_MoDa_Fields": bool(include_moda),
            "Printable_Workbook": bool(printable),
            "App_Version": "2025.11.19-p1",
        }

        c1, c2 = st.columns(2)
        with c1:
            st.subheader("Selected Villages (PPS)"); st.dataframe(df_sel_villages.sort_values(["Kebele","Selected_Order","Village"]))
        with c2:
            st.subheader("Allocation Summary"); st.dataframe(df_alloc.sort_values(["Kebele","Group","Village"]))
        st.subheader("Samples"); st.markdown("**Primaries**"); st.dataframe(df_prim.head(50)); st.markdown("**Replacements**"); st.dataframe(df_repl.head(50))

        xl_bytes = export_to_excel(df_prim, df_repl, df_sel_villages, df_alloc, settings, include_moda, printable, replacement_rate_pct)
        st.download_button(label="⬇️ Download Excel (Primaries, Replacements, Selected_Villages, Allocation_Summary, Settings_Log)", data=xl_bytes,
                           file_name=f"Sampling_Output_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        st.success("Sampling completed. Export is ready.")
        st.caption(f"Seed used: {seed} | Mode: {seed_mode} | Replacement rate: {replacement_rate_pct:.0f}%")

    except Exception as e:
        st.exception(e)
        st.error("Failed to run sampling. Please check input sheets and columns.")
